import io
import logging
from datetime import datetime
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from intelligence.profit_service import (
    compute_category_profitability,
    compute_dashboard_metrics,
    compute_employee_profitability,
    compute_product_profitability,
    compute_supplier_profitability,
)
from sales.pdf import _BOLD, _FONT, _register_fonts, _reshape

logger = logging.getLogger(__name__)

AR_FONT = 'Tahoma'
HEADER_FONT = Font(name=AR_FONT, bold=True)
BODY_FONT = Font(name=AR_FONT)
RTL_ALIGN = Alignment(horizontal='right', vertical='center')


def _header_font():
    return HEADER_FONT


def _setup_rtl_sheet(ws):
    ws.sheet_view.rightToLeft = True


def _append_row(ws, values, header=False):
    ws.append(values)
    row_idx = ws.max_row
    for cell in ws[row_idx]:
        cell.font = _header_font() if header else BODY_FONT
        cell.alignment = RTL_ALIGN


def _autosize_columns(ws):
    for col in ws.columns:
        if not col:
            continue
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 4, 10), 52)


def generate_profit_excel(store, date_from=None, date_to=None):
    wb = Workbook()
    wb.remove(wb.active)

    metrics = compute_dashboard_metrics(store)
    ws_summary = wb.create_sheet('ملخص الأرباح')
    _setup_rtl_sheet(ws_summary)
    _append_row(ws_summary, ['الفترة', 'الإيرادات', 'التكلفة', 'إجمالي الربح', 'المصاريف', 'صافي الربح'], header=True)
    for key, label in [('today', 'اليوم'), ('week', 'الأسبوع'), ('month', 'الشهر'), ('year', 'السنة')]:
        d = metrics[key]
        _append_row(ws_summary, [label, d['revenue'], d['product_cost'], d['gross_profit'], d['expenses'], d['net_profit']])
    _autosize_columns(ws_summary)

    products = compute_product_profitability(store, date_from, date_to, limit=100)
    ws_prod = wb.create_sheet('ربحية المنتجات')
    _setup_rtl_sheet(ws_prod)
    _append_row(ws_prod, ['المنتج', 'الإيرادات', 'التكلفة', 'الربح', 'الهامش %'], header=True)
    for p in products['products']:
        _append_row(ws_prod, [p['name'], p['revenue'], p['cost'], p['gross_profit'], p['margin_percent']])
    _autosize_columns(ws_prod)

    categories = compute_category_profitability(store, date_from, date_to)
    ws_cat = wb.create_sheet('ربحية التصنيفات')
    _setup_rtl_sheet(ws_cat)
    _append_row(ws_cat, ['التصنيف', 'الإيرادات', 'التكلفة', 'الربح', 'الهامش %'], header=True)
    for c in categories['categories']:
        _append_row(ws_cat, [c['name'], c['revenue'], c['cost'], c['profit'], c['margin_percent']])
    _autosize_columns(ws_cat)

    employees = compute_employee_profitability(store, date_from, date_to)
    ws_emp = wb.create_sheet('ربحية الموظفين')
    _setup_rtl_sheet(ws_emp)
    _append_row(ws_emp, ['الموظف', 'عدد الفواتير', 'الإيرادات', 'الربح', 'الهامش %'], header=True)
    for e in employees['employees']:
        _append_row(ws_emp, [e['name'], e['sales_count'], e['revenue'], e['profit'], e['margin_percent']])
    _autosize_columns(ws_emp)

    suppliers = compute_supplier_profitability(store, date_from, date_to)
    ws_sup = wb.create_sheet('تحليل الموردين')
    _setup_rtl_sheet(ws_sup)
    _append_row(ws_sup, ['المورد', 'إجمالي المشتريات', 'الكمية', 'متوسط التكلفة', 'ربح المبيعات المرتبطة'], header=True)
    for s in suppliers['suppliers']:
        _append_row(ws_sup, [s['name'], s['purchase_total'], s['items_purchased'], s['avg_unit_cost'], s['related_sales_profit']])
    _autosize_columns(ws_sup)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _content_disposition(store_name: str, ext: str) -> str:
    safe = ''.join(c if c.isalnum() or c in '-_' else '_' for c in store_name)[:40]
    date_part = datetime.now().strftime('%Y%m%d')
    ascii_name = f'profit_{safe}_{date_part}.{ext}'
    utf_name = f'تقرير_أرباح_{store_name}_{date_part}.{ext}'
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(utf_name)}'


def profit_excel_response(store, date_from=None, date_to=None):
    buffer = generate_profit_excel(store, date_from, date_to)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = _content_disposition(store.name, 'xlsx')
    return response


def _pdf_styles():
    _register_fonts()
    return {
        'title': ParagraphStyle(
            'ProfitTitle',
            fontName=_BOLD,
            alignment=TA_CENTER,
            fontSize=16,
            spaceAfter=12,
            leading=22,
        ),
        'heading': ParagraphStyle(
            'ProfitHeading',
            fontName=_BOLD,
            alignment=TA_RIGHT,
            fontSize=12,
            spaceAfter=8,
            leading=18,
        ),
        'normal': ParagraphStyle(
            'ProfitNormal',
            fontName=_FONT,
            alignment=TA_RIGHT,
            fontSize=10,
            spaceAfter=6,
            leading=16,
        ),
        'cell': ParagraphStyle(
            'ProfitCell',
            fontName=_FONT,
            alignment=TA_CENTER,
            fontSize=9,
            leading=14,
        ),
        'cell_header': ParagraphStyle(
            'ProfitCellHeader',
            fontName=_BOLD,
            alignment=TA_CENTER,
            fontSize=9,
            textColor=colors.white,
            leading=14,
        ),
    }


def _pdf_cell(text, style):
    return Paragraph(_reshape(str(text)), style)


def _pdf_table(rows, col_widths, header_rows=1):
    styles = _pdf_styles()
    table_rows = []
    for i, row in enumerate(rows):
        cell_style = styles['cell_header'] if i < header_rows else styles['cell']
        table_rows.append([_pdf_cell(cell, cell_style) for cell in row])

    table = Table(table_rows, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, header_rows - 1), colors.HexColor('#1e40af')),
        ('TEXTCOLOR', (0, 0), (-1, header_rows - 1), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, header_rows), (-1, -1), [colors.white, colors.HexColor('#f1f5f9')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    return table


def generate_profit_pdf(store, date_from=None, date_to=None):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm)
    elements = []
    styles = _pdf_styles()

    elements.append(Paragraph(_reshape(f'تقرير الأرباح — {store.name}'), styles['title']))
    elements.append(Paragraph(
        _reshape(f'تاريخ التصدير: {timezone.localtime().strftime("%Y-%m-%d %H:%M")}'),
        styles['normal'],
    ))
    elements.append(Spacer(1, 10))

    metrics = compute_dashboard_metrics(store)
    summary_data = [['الفترة', 'الإيرادات', 'صافي الربح']]
    for key, label in [('today', 'اليوم'), ('week', 'الأسبوع'), ('month', 'الشهر'), ('year', 'السنة')]:
        d = metrics[key]
        summary_data.append([label, d['revenue'], d['net_profit']])
    elements.append(_pdf_table(summary_data, [60 * mm, 50 * mm, 50 * mm]))
    elements.append(Spacer(1, 16))

    products = compute_product_profitability(store, date_from, date_to, limit=15)
    elements.append(Paragraph(_reshape('أفضل المنتجات ربحية'), styles['heading']))
    prod_data = [['المنتج', 'الإيرادات', 'الربح', 'الهامش %']]
    for p in products['products'][:15]:
        prod_data.append([p['name'], p['revenue'], p['gross_profit'], p['margin_percent']])
    elements.append(_pdf_table(prod_data, [70 * mm, 35 * mm, 35 * mm, 30 * mm]))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def profit_pdf_response(store, date_from=None, date_to=None):
    buffer = generate_profit_pdf(store, date_from, date_to)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition(store.name, 'pdf')
    return response
