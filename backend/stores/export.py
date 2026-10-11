import io
from datetime import datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from customers.models import Customer
from expenses.models import Expense
from products.models import Category, InventoryLog, Product
from purchases.models import Purchase, PurchaseItem, PurchasePayment, Supplier
from repairs.models import RepairOrder
from sales.models import Sale, SaleItem, SalePayment, SaleReturn
from stores.models import StoreMembership, StoreRolePermission


def _header_font():
    return Font(bold=True)


def _autosize_columns(ws):
    for col in ws.columns:
        if not col:
            continue
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len, 8) + 2, 48)


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _header_font()
    for row in rows:
        ws.append(row)
    _autosize_columns(ws)


def _fmt_dt(dt):
    if not dt:
        return ''
    return timezone.localtime(dt).strftime('%Y-%m-%d %H:%M')


def _fmt_user(user):
    if not user:
        return ''
    return user.get_full_name() or user.username


def generate_store_excel(store):
    wb = Workbook()
    wb.remove(wb.active)

    # معلومات المتجر
    ws_info = wb.create_sheet('معلومات المتجر')
    settings = getattr(store, 'settings', None)
    info_rows = [
        ('اسم المتجر', store.name),
        ('الهاتف', store.phone or ''),
        ('العنوان', store.address or ''),
        ('العملة', store.currency),
        ('نوع النشاط', store.get_business_type_display()),
        ('نسبة الضريبة %', str(settings.tax_rate) if settings else ''),
        ('حد تنبيه المخزون', str(settings.low_stock_threshold) if settings else ''),
        ('بادئة الفاتورة', settings.invoice_prefix if settings else ''),
        ('تاريخ التصدير', timezone.localtime().strftime('%Y-%m-%d %H:%M')),
    ]
    _write_sheet(ws_info, ['الحقل', 'القيمة'], info_rows)

    # التصنيفات
    ws_cat = wb.create_sheet('التصنيفات')
    _write_sheet(
        ws_cat,
        ['الاسم', 'الوصف'],
        [(c.name, c.description) for c in Category.objects.filter(store=store).order_by('name')],
    )

    # المنتجات
    ws_prod = wb.create_sheet('المنتجات')
    _write_sheet(
        ws_prod,
        ['الاسم', 'SKU', 'باركود', 'IMEI', 'التصنيف', 'سعر الشراء', 'سعر البيع', 'الكمية', 'حد التنبيه', 'نشط'],
        [
            (
                p.name,
                p.sku,
                p.barcode,
                p.imei,
                p.category.name if p.category_id else '',
                float(p.buy_price),
                float(p.sell_price),
                p.quantity,
                p.low_stock_threshold,
                'نعم' if p.is_active else 'لا',
            )
            for p in Product.objects.filter(store=store).select_related('category').order_by('name')
        ],
    )

    # العملاء
    ws_cust = wb.create_sheet('العملاء')
    _write_sheet(
        ws_cust,
        ['الاسم', 'الجوال', 'البريد', 'العنوان', 'الدين', 'نشط'],
        [
            (
                c.name,
                c.phone,
                c.email,
                c.address,
                float(c.debt),
                'نعم' if c.is_active else 'لا',
            )
            for c in Customer.objects.filter(store=store).order_by('name')
        ],
    )

    # الفواتير
    ws_sales = wb.create_sheet('الفواتير')
    _write_sheet(
        ws_sales,
        [
            'رقم الفاتورة', 'العميل', 'الحالة', 'طريقة الدفع', 'المجموع', 'الخصم',
            'الضريبة', 'الإجمالي', 'المدفوع', 'المتبقي', 'المرتجعات', 'ملاحظات', 'التاريخ',
        ],
        [
            (
                s.invoice_number,
                s.customer.name if s.customer_id else '',
                s.get_status_display(),
                s.get_payment_method_display(),
                float(s.subtotal),
                float(s.discount),
                float(s.tax_amount),
                float(s.total),
                float(s.amount_paid),
                float(max(0, s.total - s.total_returned - s.amount_paid)),
                float(s.total_returned),
                s.notes,
                _fmt_dt(s.created_at),
            )
            for s in Sale.objects.filter(store=store).select_related('customer').order_by('-created_at')
        ],
    )

    # بنود الفواتير
    ws_items = wb.create_sheet('بنود الفواتير')
    _write_sheet(
        ws_items,
        ['رقم الفاتورة', 'المنتج', 'SKU', 'IMEI', 'الكمية', 'السعر', 'الإجمالي'],
        [
            (
                item.sale.invoice_number,
                item.product_name,
                item.sku,
                item.imei,
                item.quantity,
                float(item.unit_price),
                float(item.line_total),
            )
            for item in SaleItem.objects.filter(sale__store=store).select_related('sale').order_by('-sale__created_at')
        ],
    )

    # دفعات الفواتير (إكمال الدفع الجزئي)
    ws_sale_pay = wb.create_sheet('دفعات الفواتير')
    _write_sheet(
        ws_sale_pay,
        ['رقم الفاتورة', 'المبلغ', 'طريقة الدفع', 'ملاحظات', 'بواسطة', 'التاريخ'],
        [
            (
                p.sale.invoice_number,
                float(p.amount),
                p.get_method_display(),
                p.notes,
                _fmt_user(p.created_by),
                _fmt_dt(p.created_at),
            )
            for p in SalePayment.objects.filter(store=store).select_related('sale', 'created_by').order_by('-created_at')
        ],
    )

    # مرتجعات الفواتير
    ws_returns = wb.create_sheet('مرتجعات الفواتير')
    _write_sheet(
        ws_returns,
        ['رقم المرتجع', 'رقم الفاتورة', 'مبلغ الاسترداد', 'ملاحظات', 'التاريخ'],
        [
            (
                r.return_number,
                r.sale.invoice_number,
                float(r.refund_amount),
                r.notes,
                _fmt_dt(r.created_at),
            )
            for r in SaleReturn.objects.filter(store=store).select_related('sale').order_by('-created_at')
        ],
    )

    # الموردون
    ws_suppliers = wb.create_sheet('الموردون')
    _write_sheet(
        ws_suppliers,
        ['الاسم', 'الجوال', 'البريد', 'العنوان', 'المستحقات', 'ملاحظات', 'نشط'],
        [
            (
                s.name,
                s.phone,
                s.email,
                s.address,
                float(s.balance),
                s.notes,
                'نعم' if s.is_active else 'لا',
            )
            for s in Supplier.objects.filter(store=store).order_by('name')
        ],
    )

    # فواتير الشراء
    ws_purchases = wb.create_sheet('فواتير الشراء')
    _write_sheet(
        ws_purchases,
        [
            'رقم الفاتورة', 'المورد', 'الحالة', 'طريقة الدفع', 'المجموع', 'الخصم',
            'الضريبة', 'الإجمالي', 'المدفوع', 'المتبقي', 'ملاحظات', 'بواسطة', 'التاريخ',
        ],
        [
            (
                p.invoice_number,
                p.supplier.name if p.supplier_id else '',
                p.get_status_display(),
                p.get_payment_method_display(),
                float(p.subtotal),
                float(p.discount),
                float(p.tax_amount),
                float(p.total),
                float(p.amount_paid),
                float(max(0, p.total - p.amount_paid)),
                p.notes,
                _fmt_user(p.created_by),
                _fmt_dt(p.created_at),
            )
            for p in Purchase.objects.filter(store=store).select_related('supplier', 'created_by').order_by('-created_at')
        ],
    )

    # بنود فواتير الشراء
    ws_purchase_items = wb.create_sheet('بنود فواتير الشراء')
    _write_sheet(
        ws_purchase_items,
        ['رقم الفاتورة', 'المنتج', 'SKU', 'الكمية', 'سعر الشراء', 'الإجمالي'],
        [
            (
                item.purchase.invoice_number,
                item.product_name,
                item.sku,
                item.quantity,
                float(item.unit_cost),
                float(item.line_total),
            )
            for item in PurchaseItem.objects.filter(purchase__store=store).select_related('purchase').order_by('-purchase__created_at')
        ],
    )

    # دفعات فواتير الشراء
    ws_purchase_pay = wb.create_sheet('دفعات فواتير الشراء')
    _write_sheet(
        ws_purchase_pay,
        ['رقم الفاتورة', 'المبلغ', 'طريقة الدفع', 'ملاحظات', 'بواسطة', 'التاريخ'],
        [
            (
                p.purchase.invoice_number,
                float(p.amount),
                p.get_method_display(),
                p.notes,
                _fmt_user(p.created_by),
                _fmt_dt(p.created_at),
            )
            for p in PurchasePayment.objects.filter(store=store).select_related('purchase', 'created_by').order_by('-created_at')
        ],
    )

    # سجل المخزون
    ws_inv_log = wb.create_sheet('سجل المخزون')
    _write_sheet(
        ws_inv_log,
        ['المنتج', 'العملية', 'التغيير', 'قبل', 'بعد', 'المرجع', 'ملاحظات', 'بواسطة', 'التاريخ'],
        [
            (
                log.product.name,
                log.get_action_display(),
                log.quantity_change,
                log.quantity_before,
                log.quantity_after,
                log.reference,
                log.notes,
                _fmt_user(log.created_by),
                _fmt_dt(log.created_at),
            )
            for log in InventoryLog.objects.filter(store=store).select_related('product', 'created_by').order_by('-created_at')
        ],
    )

    # المصاريف
    ws_exp = wb.create_sheet('المصاريف')
    _write_sheet(
        ws_exp,
        ['العنوان', 'المبلغ', 'التصنيف', 'طريقة الدفع', 'التاريخ', 'ملاحظات'],
        [
            (
                e.title,
                float(e.amount),
                e.get_category_display(),
                e.get_payment_method_display(),
                e.expense_date.isoformat(),
                e.notes,
            )
            for e in Expense.objects.filter(store=store, is_active=True).order_by('-expense_date')
        ],
    )

    # الصيانة
    ws_rep = wb.create_sheet('الصيانة')
    _write_sheet(
        ws_rep,
        ['رقم التذكرة', 'العميل', 'الجوال', 'الجهاز', 'IMEI', 'الحالة', 'التكلفة', 'العربون', 'تاريخ الاستلام'],
        [
            (
                r.ticket_number,
                r.customer_name or (r.customer.name if r.customer_id else ''),
                r.customer_phone,
                f'{r.device_brand} {r.device_model}'.strip(),
                r.imei,
                r.get_status_display(),
                float(r.total_cost),
                float(r.deposit_paid),
                timezone.localtime(r.received_at).strftime('%Y-%m-%d %H:%M') if r.received_at else '',
            )
            for r in RepairOrder.objects.filter(store=store).select_related('customer').order_by('-created_at')
        ],
    )

    # الفريق
    ws_team = wb.create_sheet('الفريق')
    _write_sheet(
        ws_team,
        ['الاسم', 'البريد', 'الدور', 'نشط'],
        [
            (
                m.user.get_full_name() or m.user.username,
                m.user.email,
                m.role.name,
                'نعم' if m.is_active else 'لا',
            )
            for m in StoreMembership.objects.filter(store=store).select_related('user', 'role').order_by('joined_at')
        ],
    )

    # صلاحيات الأدوار
    ws_perms = wb.create_sheet('صلاحيات الأدوار')
    _write_sheet(
        ws_perms,
        ['الدور', 'الصلاحيات'],
        [
            (
                rp.role.name,
                ', '.join(str(p) for p in (rp.permissions or [])),
            )
            for rp in StoreRolePermission.objects.filter(store=store).select_related('role').order_by('role__name')
        ],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def excel_http_response(store):
    buffer = generate_store_excel(store)
    safe_name = ''.join(c if c.isalnum() or c in '-_' else '_' for c in store.name)[:40]
    stamp = datetime.now().strftime('%Y%m%d')
    filename = f'mobishop_{safe_name}_{stamp}.xlsx'
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Content-Length'] = len(response.content)
    return response
